#!/usr/bin/env python3
"""Импорт таблицы оператора «Судебная практика» в evidence/practice.

  source/sudebnaya-praktika-2026-10-13.xlsx → cases.tsv (лист «Все дела» + practice_id, norm_ids, scenario, source_class),
                                              not-found.tsv (лист «Практика не найдена»), practice.json (исходные колонки).
Привязка norm_ids — по регулярным выражениям над колонкой «Норма (закон, статья)» (RULES);
сценарий — по разделу и ключевым словам колонок «К чему относится» и «Описание» (scenario()).
Это эвристика для навигации, не правовой вывод: строки остаются наводками (evidence/CLAUDE.md §0.2).
Запуск: python3 -B scripts/import_practice_xlsx.py  (нужен openpyxl); затем scripts/build_practice_registry.py
"""
import json,re,csv,collections,sys
from pathlib import Path
import openpyxl
PR = Path(__file__).resolve().parents[1] / "evidence" / "practice"
XLSX = PR / "source" / "sudebnaya-praktika-2026-10-13.xlsx"
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["Все дела"]; hdr = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
json.dump(rows, open(PR / "practice.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
nf = [r for r in wb["Практика не найдена"].iter_rows(values_only=True) if r[0] is not None]
with open(PR / "not-found.tsv", "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f, delimiter="\t"); w.writerow(["n", "norm_text", "ui_aspect", "note"])
    for r in nf[1:]: w.writerow([int(float(r[0])), r[1], r[2], r[3]])
# (regex on norm string, [norm ids])  — order matters; all matching rules apply
RULES=[
 (r'149-ФЗ ст\. 10\.6', ['ru-social-network-duties']),
 (r'149-ФЗ ст\. 10\.2-2|КоАП 13\.56|КоАП ст\. 13\.56', ['ru-149fz-10-2-2','ru-koap-13-56']),
 (r'КоАП 13\.55|406-ФЗ', ['ru-149fz-10-auth','ru-koap-13-55']),
 (r'149-ФЗ ст\. 10 ч\. 2|149-ФЗ ст\. 10(?![.\d])', ['ru-149fz-10-2']),
 (r'149-ФЗ ст\. 15\.8|VPN', ['ru-vpn-promotion-ban']),
 (r'152-ФЗ ст\. (6|9)\b|152-ФЗ ст\. 9;|152-ФЗ; КоАП 13\.11 ч\. 1|152-ФЗ ст\. 6, 9', ['ru-152fz-9']),
 (r'152-ФЗ ст\. 15\b', ['ru-152fz-15']),
 (r'152-ФЗ ст\. 18\.1', ['ru-152fz-18-1']),
 (r'152-ФЗ ст\. (14|20|21)\b', ['ru-152fz-withdrawal']),
 (r'152-ФЗ ст\. 16\b', ['ru-152fz-16']),
 (r'152-ФЗ ст\. 11\b|572-ФЗ|биометр', ['ru-152fz-11-biometric']),
 (r'152-ФЗ ст\. 19|ФСТЭК|ПП 1119|152-ФЗ ст\. 21', ['ru-fstek-ui-measures']),
 (r'КоАП 13\.11\b(?!\.)', ['ru-koap-13-11']),
 (r'353-ФЗ ст\. 7|353-ФЗ ст\. 11', ['ru-353fz-7-addons-cooling']),
 (r'353-ФЗ ст\. 6', ['ru-353fz-psk-frame']),
 (r'259-ФЗ', ['ru-259fz-no-crypto']),
 (r'ЗоЗПП ст\. 10\.1', ['ru-interface-language']),
 (r'ЗоЗПП ст\. (8|9|10)\b|ЗоЗПП ст\. 8[–-]1[02]|ЗоЗПП ст\. 10,', ['ru-zozpp-10']),
 (r'ЗоЗПП ст\. 12\b|ЗоЗПП ст\. 8[–-]12', ['ru-zozpp-12']),
 (r'ЗоЗПП ст\. 16\b(?!\.)', ['ru-zozpp-16']),
 (r'ЗоЗПП ст\. 16\.1', ['ru-zozpp-16-1']),
 (r'ЗоЗПП ст\. 32|ГК 782', ['ru-zozpp-32']),
 (r'ЗоЗПП ст\. 26\.1|ПП 2463', ['ru-zozpp-26-1']),
 (r'ЗоЗПП ст\. 46', ['ru-zozpp-46']),
 (r'КоАП 14\.8', ['ru-koap-14-8']),
 (r'289-ФЗ', ['ru-289fz-platform']),
 (r'54-ФЗ|КоАП 14\.5', ['ru-54fz-receipt','ru-koap-14-5']),
 (r'ГК 426', ['ru-gk-426']),
 (r'ГК (428|435|437|438|494|497)\b|оферта', ['ru-gk-oferta']),
 (r'38-ФЗ ст\. 5\b(?!\.)|38-ФЗ ст\. 5;|38-ФЗ ст\. 5 ч', ['ru-38fz-5-3']),
 (r'38-ФЗ ст\. 18\b(?!\.)', ['ru-38fz-18']),
 (r'38-ФЗ ст\. 18\.1|38-ФЗ ст\. 5 ч\. 9|38-ФЗ ст\. 8\b', ['ru-38fz-18-1']),
 (r'38-ФЗ ст\. 9\b', ['ru-38fz-9']),
 (r'38-ФЗ ст\. 6\b', ['ru-38fz-6-child-ads']),
 (r'38-ФЗ ст\. 5 ч\. 6|религиоз', ['ru-religious-symbols-promo']),
 (r'КоАП 14\.3', ['ru-koap-14-3']),
 (r'КоАП 6\.13', ['ru-koap-6-13','ru-drug-symbols']),
 (r'КоАП 6\.21', ['ru-koap-6-21','ru-lgbt-symbolism']),
 (r'КоАП 20\.3\b(?!\d)', ['ru-koap-20-3','ru-extremist-symbols']),
 (r'КоАП 20\.29', ['ru-koap-20-29']),
 (r'КоАП 20\.33|272-ФЗ', ['ru-koap-20-33','ru-undesirable-orgs']),
 (r'КоАП 19\.34|255-ФЗ|иноагент', ['ru-koap-19-34-1','ru-foreign-agent-labeling']),
 (r'КоАП 17\.10|ФКЗ', ['ru-koap-17-10','ru-state-symbols']),
 (r'КоАП 6\.17|436-ФЗ ст\. 12', ['ru-koap-6-17','ru-age-marking']),
 (r'436-ФЗ ст\. 5|436-ФЗ ст\. 14', ['ru-436fz-child-sections']),
 (r'КоАП 5\.26', ['ru-koap-5-26']),
 (r'ГК 152\.1', ['ru-gk-152-1-likeness']),
 (r'135-ФЗ ст\. 14\.6|гл\. 2\.1 135-ФЗ', ['ru-135fz-style-imitation']),
 (r'КоАП 14\.33', ['ru-koap-14-33']),
 (r'181-ФЗ|ГОСТ Р 52872|ПП 102|(?<![\d])8-ФЗ', ['ru-pp-102-a11y','ru-gost-52872','ru-181fz-15-captions']),
 (r'236-ФЗ', ['ru-236fz-landing']),
 (r'ПП 236|210-ФЗ', ['ru-pp-236-eforms']),
 (r'162-ФЗ', ['ru-162fz-26']),
 (r'закон об ИИ|\bИИ\b', ['ru-ai-generation-rights']),
 (r'2124-1|СМИ', ['ru-media-imprint']),
 (r'3854-У', ['ru-cbr-3854u-cooling']),
 (r'ФАП|Минтранс', ['ru-fap-82-air-addons']),
 (r'Указ|Банк России', []),
]
# scenario by section + keywords
def scenario(r):
    sec=r['Раздел']; what=(str(r['К чему относится в интерфейсе'])+' '+str(r['Описание нарушения'])).lower(); norm=str(r['Норма (закон, статья)'])
    if sec.startswith('5.'):
        if '135-ФЗ' in norm or 'оферта площадки' in norm or 'штраф продавцу' in what or 'штраф за' in what: return 'platform'
        if 'подписк' in what or 'автопродл' in what or 'списан' in what: return 'G4'
        if 'ст. 32' in norm or 'отказ от услуг' in what or 'отмен' in what: return 'G5'
        if 'сбор' in what or 'стоимост' in what or 'цена' in what: return 'G2/G3'
        if 'галочк' in what or 'предустанов' in what or 'автоматически' in what or 'страхов' in what: return 'G1'
        if 'сведени' in what or 'продавц' in what or 'исполнител' in what or 'карточк' in what: return 'G3'
        return 'G2/G3'
    if sec.startswith('6.'):
        if '10.7' in norm or 'заблокирован' in what or 'ограниченным доступом' in what or 'vpn' in what or 'обхода блокировок' in norm: return 'content'
        if 'рассылк' in what or 'сообщени' in what or 'согласие на рекламу' in what: return 'G7'
        if 'erid' in what or 'маркиров' in what or 'токен' in what or 'пометка' in what or 'сведения о рекламе' in what or 'отчётност' in what: return 'G8'
        return 'G6'
    if sec.startswith(('7.','9.','8.','1.','11.')): return 'content'
    if sec.startswith('2.'): return 'consent'
    if sec.startswith('3.'): return 'security'
    if sec.startswith('4.'): return 'вне-охвата:finance'
    if sec.startswith('10.'): return 'вне-охвата:ip'
    if sec.startswith('12.'): return 'вне-охвата:hr'
    if sec.startswith('13.'): return 'вне-охвата:gov'
    if sec.startswith('14.'): return 'вне-охвата:antitrust'
    return 'вне-охвата'
src_map={'текст акта или карточка суда':'act_or_card','новости или обзор':'news','только поисковая выдача':'search'}
out=[]; unmatched=collections.Counter(); gaps=collections.Counter()
for r in rows:
    n=str(r['Норма (закон, статья)']); ids=[]
    for rx,v in RULES:
        if re.search(rx,n): ids+= [i for i in v if i not in ids]
    if not ids: unmatched[n]+=1
    # norms named but absent from KB
    for tok in re.findall(r'(?:КоАП|УК|ГК|ТК) ?(?:ст\. )?[\d.]+(?: ч\. [\d.]+)?|\d+-ФЗ(?: ст\. [\d.–-]+)?',n):
        if not any(re.search(rx,tok) for rx,_ in RULES if _): gaps[tok]+=1
    out.append({'practice_id':f"P-{int(float(r['№'])):03d}", 'section':r['Раздел'], 'norm_text':n, 'norm_ids':';'.join(ids), 'scenario':scenario(r),
        'ui_aspect':r['К чему относится в интерфейсе'], 'case':r['Дело / номер'], 'authority_date':r['Суд или орган, дата'], 'parties':r['Кто привлечён / стороны'],
        'violation':r['Описание нарушения'], 'outcome':r['Решение — итог'], 'url':r['Ссылка'], 'source_class':src_map.get(r['Источник сведений'],r['Источник сведений'])})
with open(PR / 'cases.tsv','w',encoding='utf-8',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(out[0].keys()),delimiter='\t'); w.writeheader(); w.writerows(out)
print("rows",len(out),"mapped",sum(1 for o in out if o['norm_ids']),"unmapped",sum(1 for o in out if not o['norm_ids']))
print("\n=== unmatched norm strings ==="); [print(f"{v:3} {k}") for k,v in unmatched.most_common()]
print("\n=== norm tokens absent from KB (counted) ==="); [print(f"{v:3} {k}") for k,v in gaps.most_common(40)]
print("\n=== scenario distribution ==="); print(collections.Counter(o['scenario'] for o in out))
print("\n=== norm_ids distribution ==="); print(collections.Counter(i for o in out for i in o['norm_ids'].split(';') if i).most_common())
